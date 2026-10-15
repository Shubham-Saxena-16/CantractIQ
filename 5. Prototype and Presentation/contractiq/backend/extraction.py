"""Convert uploaded PDF / DOCX / Markdown / text contracts into Markdown, preserving headings,
clause numbers, tables and page markers where the source allows it.

Errors are reported, never swallowed: an unreadable page or file produces an explicit warning or
ExtractionError rather than an empty "clean" result.
"""
from __future__ import annotations

import io
import re
import shutil


class ExtractionError(Exception):
    pass


def extract(filename: str, data: bytes) -> tuple[str, list[dict], dict]:
    """Return (markdown, warnings, info)."""
    name = filename.lower()
    if not data:
        raise ExtractionError("The file is empty.")
    if name.endswith(".pdf"):
        return extract_pdf(data)
    if name.endswith(".docx"):
        return extract_docx(data)
    if name.endswith((".md", ".markdown", ".txt")):
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
            return text, [{"level": "warning", "code": "encoding", "message": "File was not UTF-8; decoded as Latin-1."}], {"format": "text"}
        return text, [], {"format": "markdown" if name.endswith((".md", ".markdown")) else "text"}
    if name.endswith(".doc"):
        raise ExtractionError("Legacy .doc files are not supported. Save the document as .docx or PDF and upload again.")
    raise ExtractionError("Unsupported file type. Upload a PDF, DOCX, Markdown or text file.")


# ---------------------------------------------------------------- PDF

HEADING_LINE = re.compile(r"^(ARTICLE\s+\d+|SCHEDULE\s+\w+|EXHIBIT\s+\w+|ANNEX(URE)?\s+\w+|§\s*\d+|\d{1,2}\.?\s+[A-Z][A-Z &,'()/-]{3,})(\s*[—:-].*)?$")


def _ocr_available() -> tuple[bool, str]:
    try:
        import pytesseract  # noqa: F401
        import pypdfium2  # noqa: F401
    except ImportError:
        return False, "OCR libraries not installed (pip install pytesseract pypdfium2)"
    if not shutil.which("tesseract"):
        return False, "Tesseract OCR binary not found on PATH"
    return True, ""


def _ocr_page(data: bytes, index: int) -> str:
    import pypdfium2 as pdfium
    import pytesseract
    pdf = pdfium.PdfDocument(data)
    img = pdf[index].render(scale=300 / 72).to_pil()
    return pytesseract.image_to_string(img)


def _lines_to_md(text: str) -> str:
    out = []
    for line in text.splitlines():
        s = line.strip()
        if not s:
            out.append("")
            continue
        if len(s) < 90 and HEADING_LINE.match(s) and not re.match(r"^\d+\.\d+", s):
            out.append(f"\n## {s}\n")
        else:
            out.append(s)
    md = "\n".join(out)
    return re.sub(r"\n{3,}", "\n\n", md)


def extract_pdf(data: bytes) -> tuple[str, list[dict], dict]:
    try:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError
    except ImportError as e:  # pragma: no cover
        raise ExtractionError(f"PDF support unavailable: {e}")
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as e:
        raise ExtractionError(f"Could not open the PDF (it may be corrupt or not a PDF): {e}")
    if reader.is_encrypted:
        try:
            if reader.decrypt("") == 0:
                raise ExtractionError("The PDF is password-protected. Remove the password and upload again.")
        except Exception:
            raise ExtractionError("The PDF is password-protected. Remove the password and upload again.")
    warnings, parts = [], []
    ocr_ok, ocr_reason = _ocr_available()
    unreadable, ocred = [], []
    n = len(reader.pages)
    for i, page in enumerate(reader.pages, start=1):
        try:
            txt = page.extract_text() or ""
        except Exception as e:
            txt = ""
            warnings.append({"level": "error", "code": "page_error", "message": f"Page {i}: text extraction failed ({e})."})
        if len(txt.strip()) < 25:
            if ocr_ok:
                try:
                    txt = _ocr_page(data, i - 1)
                    ocred.append(i)
                except Exception as e:
                    warnings.append({"level": "error", "code": "ocr_failed", "message": f"Page {i}: OCR failed ({e})."})
            if len(txt.strip()) < 25:
                unreadable.append(i)
        parts.append(f"<!-- page {i} -->\n{_lines_to_md(txt)}")
    if ocred:
        warnings.append({"level": "warning", "code": "ocr_used",
                         "message": f"OCR was used on page(s) {', '.join(map(str, ocred))}. Verify numbers and names against the original."})
    if unreadable:
        warnings.append({"level": "error", "code": "unreadable_pages",
                         "message": f"No readable text on page(s) {', '.join(map(str, unreadable))} of {n} — likely scanned images. "
                                    + ("" if ocr_ok else f"OCR is not available: {ocr_reason}. ")
                                    + "Clauses on these pages were NOT analysed; do not treat them as absent."})
    md = "\n\n".join(parts)
    if not re.search(r"[A-Za-z]{4}", md.replace("page", "")):
        raise ExtractionError("No text could be extracted from this PDF. It appears to be scanned and OCR is "
                              f"unavailable ({ocr_reason or 'OCR produced no text'}).")
    return md, warnings, {"format": "pdf", "pages": n, "ocr_pages": ocred, "unreadable_pages": unreadable}


# ---------------------------------------------------------------- DOCX

def extract_docx(data: bytes) -> tuple[str, list[dict], dict]:
    try:
        import docx
        from docx.document import Document as _D
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as e:  # pragma: no cover
        raise ExtractionError(f"DOCX support unavailable: {e}")
    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise ExtractionError(f"Could not open the DOCX file (it may be corrupt or not a Word document): {e}")
    warnings, lines = [], []
    auto_numbered = 0
    body = d.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            p = Paragraph(child, d)
            text = p.text.strip()
            if not text:
                lines.append("")
                continue
            style = (p.style.name if p.style is not None else "") or ""
            if child.find(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numPr") is not None:
                auto_numbered += 1
            m = re.match(r"Heading (\d)", style)
            if style == "Title":
                lines.append(f"# {text}")
            elif m:
                lines.append("#" * min(int(m.group(1)) + 1, 4) + " " + text)
            elif len(text) < 80 and all(r.bold for r in p.runs if r.text.strip()) and p.runs and text.upper() == text:
                lines.append(f"## {text}")
            else:
                lines.append(text)
            lines.append("")
        elif tag == "tbl":
            t = Table(child, d)
            rows = []
            for r in t.rows:
                cells = [c.text.strip().replace("\n", " ").replace("|", "/") for c in r.cells]
                rows.append("| " + " | ".join(cells) + " |")
            if rows:
                ncols = rows[0].count("|") - 1
                rows.insert(1, "|" + "---|" * ncols)
                lines.extend(rows + [""])
    if auto_numbered:
        warnings.append({"level": "warning", "code": "auto_numbering",
                         "message": f"{auto_numbered} paragraph(s) use Word automatic numbering, which is not stored as text. "
                                    "Clause numbers for those paragraphs may be missing from citations."})
    has_headers = any(any(p.text.strip() for p in s.header.paragraphs + s.footer.paragraphs) for s in d.sections)
    if has_headers:
        warnings.append({"level": "info", "code": "headers_skipped", "message": "Page headers/footers were not included."})
    md = re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()
    if len(md) < 50:
        raise ExtractionError("The DOCX file contains almost no text (it may contain only images).")
    return md, warnings, {"format": "docx", "paragraphs": len(d.paragraphs), "tables": len(d.tables)}
